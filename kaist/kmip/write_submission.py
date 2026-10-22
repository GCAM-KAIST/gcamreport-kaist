"""Write the KMIP submission workbook from the step4 result.

Layout follows the submitted file (sheet "Template(제출용)"): the main rows of
every scenario one block after another, then the "Added" rows of every
scenario. Cell styles and the 비고 notes are copied from the template block.
Missing values are written as 0, as in the past GC submissions.

Usage: python write_submission.py TEMPLATE.xlsx STEP4_RESULT.csv OUT.xlsx [MODEL]
Called at the end of step4_fill_template.R. Needs pandas and openpyxl.
"""
import sys
from copy import copy

import openpyxl
import pandas as pd

SHEET = "Template(제출용)"


def write_submission(template_path, csv_path, out_path, model="GC"):
    df = pd.read_csv(csv_path)
    wb = openpyxl.load_workbook(template_path)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    year_cols = [i for i, h in enumerate(header) if isinstance(h, int)]
    col = {h: i for i, h in enumerate(header)}

    # template block: values and styles of each row
    block = []
    for r in range(2, ws.max_row + 1):
        cells = list(ws[r])
        block.append(([c.value for c in cells],
                      [(copy(c.font), copy(c.fill), copy(c.border), copy(c.alignment), c.number_format)
                       for c in cells]))
    main = [b for b in block if b[0][col["Region"]] != "Added"]
    added = [b for b in block if b[0][col["Region"]] == "Added"]

    scens = list(dict.fromkeys(df["Scenario"]))
    vals = df.drop_duplicates(["Scenario", "Variable"]).set_index(["Scenario", "Variable"])
    rows = [(s, b) for s in scens for b in main] + [(s, b) for s in scens for b in added]

    ws.delete_rows(2, ws.max_row)
    for i, (scen, (values, styles)) in enumerate(rows, start=2):
        values = list(values)
        values[col["model"]] = model
        values[col["scenario"]] = scen
        var = values[col["variable"]]
        for j in year_cols:
            key = (scen, var)
            v = vals.at[key, str(header[j])] if key in vals.index else None
            values[j] = 0.0 if v is None or pd.isna(v) else float(v)
        for j, (v, st) in enumerate(zip(values, styles), start=1):
            cell = ws.cell(i, j, v)
            cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = st
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(header))}{len(rows) + 1}"
    wb.save(out_path)
    print(f"KMIP submission workbook: {out_path} ({len(rows)} rows, {len(scens)} scenarios)")


if __name__ == "__main__":
    write_submission(*sys.argv[1:])
